"""Offline course filing. No cloud credentials or AI calls are used."""
import argparse
import hashlib
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
CONFIG = ROOT / 'local' / 'schedule.json'
SUPPORTED = {'.txt', '.md', '.docx', '.pdf'}


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temp.replace(path)


def today():
    # Recording dates use Beijing time, matching Feishu exports.
    from datetime import timezone, timedelta
    return datetime.now(timezone(timedelta(hours=8))).date().isoformat()


def candidates(cfg, day):
    d = date.fromisoformat(day)
    override = cfg.get('date_overrides', {}).get(day, {})
    weekday = override.get('weekday', d.isoweekday())
    week = override.get('week')
    if week is None and cfg.get('first_week_monday'):
        week = (d - date.fromisoformat(cfg['first_week_monday'])).days // 7 + 1
    found = []
    enabled = {c['id'] for c in cfg['courses'] if c.get('archive_enabled', True)}
    for lesson in cfg['lessons']:
        if lesson['weekday'] != weekday or lesson['course'] not in enabled:
            continue
        weeks = lesson.get('weeks', []) + [w for a, b in lesson.get('week_ranges', []) for w in range(a, b + 1)]
        if week is None or week in weeks:
            if lesson not in found:
                found.append(dict(lesson))
    return {'date': day, 'weekday': weekday, 'week': week, 'note': override.get('note', ''), 'lessons': found}


def extract(path):
    if path.suffix.lower() == '.pdf':
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            value = '\n\n'.join(p.extract_text() or '' for p in pdf.pages)
    elif path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(path)
        value = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    else:
        raw = path.read_bytes()
        for enc in ('utf-8-sig', 'utf-16', 'gb18030'):
            try:
                value = raw.decode(enc)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('无法识别文本编码')
    if not value.strip():
        raise ValueError('没有可提取文字；扫描 PDF 需先识别文字')
    return value.replace('\r\n', '\n').replace('\r', '\n')


def identify(cfg, title, body):
    # Longest matching alias prevents 普通化学 from matching 普通化学实验.
    def matches(text):
        hits = []
        for c in cfg['courses']:
            for alias in c['aliases']:
                if alias.casefold() in text.casefold():
                    hits.append((alias, c['id']))
        return {cid for alias, cid in hits if not any(alias != other and alias.casefold() in other.casefold() for other, _ in hits)}
    hits = matches(title)
    method = '文件标题'
    if not hits:
        hits = matches(body)
        method = '正文课程名'
    if len(hits) == 1:
        return next(iter(hits)), method
    return None, '出现多个课程名称' if hits else '没有明确课程名称'


def safe_name(text):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', text).strip(' .')[:70] or '课堂资料'


def archive_stem(course, day, body):
    """Readable chronological names; IDs and hashes stay in the private index."""
    stamp = re.search(r'录音开始：' + re.escape(day) + r' (\d{2}):(\d{2}):(\d{2})', body)
    clock = '_'+ '-'.join(stamp.groups()) if stamp else ''
    if re.search(r'^资料类型：飞书AI总结$', body, re.M):
        kind = '飞书AI总结'
    elif re.search(r'^资料类型：飞书逐字稿', body, re.M):
        kind = '飞书逐字稿'
    else:
        kind = '课堂资料'
    return f'{day}{clock}_{safe_name(course["name"])}_{kind}'


def recording_start(day, body):
    match = re.search(r'录音开始：' + re.escape(day) + r' (\d{2}:\d{2}:\d{2})', body)
    if match:
        value = day+' '+match.group(1)
        datetime.strptime(value, '%Y-%m-%d %H:%M:%S')
        return value
    return None


def allocate_paths(folder, day, stem, suffix, layout='flat'):
    """An incremental suffix preserves distinct exports without overwriting."""
    serial = 1
    if layout not in {'flat', 'lesson_folders'}:
        raise ValueError('archive_layout 应为 lesson_folders 或 flat')
    while True:
        if layout == 'lesson_folders':
            session = re.sub(r'_(?:飞书AI总结|飞书逐字稿|课堂资料)$', '', stem)
            session = session if serial == 1 else f'{session}_{serial:02d}'
            lesson_folder = folder / '课堂归档' / day[:7] / session
            if not lesson_folder.exists():
                label = '课堂总结' if stem.endswith('_飞书AI总结') else ('课堂逐字稿' if stem.endswith('_飞书逐字稿') else '课堂资料')
                return (lesson_folder / (label+'.pdf'), lesson_folder / '原始资料' / ('导出原文'+suffix),
                        lesson_folder / '原始资料' / '提取文字.txt')
            serial += 1
            continue
        candidate = stem if serial == 1 else f'{stem}_{serial:02d}'
        paths = (folder / '课堂归档' / day[:7] / (candidate+'.pdf'),
                 folder / '原始资料' / day[:7] / (candidate+suffix),
                 folder / '原始资料' / day[:7] / (candidate+'.source.txt'))
        if not any(p.exists() for p in paths):
            return paths
        serial += 1


def course_folder(cfg, course):
    desktop = Path(cfg['desktop']).resolve()
    folder = (desktop / course['folder']).resolve()
    if folder == desktop or not folder.is_relative_to(desktop):
        raise ValueError('课程文件夹必须位于配置的桌面目录内')
    return folder


def source_day(path, fallback):
    found = re.search(r'(?<!\d)(20\d{2})[-_.年](\d{1,2})[-_.月](\d{1,2})(?:日)?(?!\d)', path.stem)
    if found:
        return date(*map(int, found.groups())).isoformat(), '文件名日期'
    if fallback:
        return date.fromisoformat(fallback).isoformat(), '用户指定批次上课日期'
    raise ValueError('文件名无日期，请用 --date 指定上课日期；不采用文件修改时间')


def write_pdf(target, course, day, source, body):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font = 'CourseChinese'
    if font not in pdfmetrics.getRegisteredFontNames():
        fontpath = Path('C:/Windows/Fonts/msyh.ttc')
        if not fontpath.is_file():
            raise ValueError('缺少中文字体 C:/Windows/Fonts/msyh.ttc，请配置本机字体')
        pdfmetrics.registerFont(TTFont(font, str(fontpath), subfontIndex=0))
        bold_path = Path('C:/Windows/Fonts/msyhbd.ttc')
        if bold_path.is_file():
            pdfmetrics.registerFont(TTFont(font + 'Bold', str(bold_path), subfontIndex=0))
            pdfmetrics.registerFontFamily(font, normal=font, bold=font + 'Bold', italic=font, boldItalic=font + 'Bold')
        else:
            pdfmetrics.registerFontFamily(font, normal=font, bold=font, italic=font, boldItalic=font)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_suffix('.tmp.pdf')
    heading = ParagraphStyle('heading', fontName=font, fontSize=17, leading=25, spaceAfter=14)
    small = ParagraphStyle('small', fontName=font, fontSize=9, leading=15, textColor='#555555')
    normal = ParagraphStyle('normal', fontName=font, fontSize=10, leading=17, spaceAfter=5, wordWrap='CJK')
    section = ParagraphStyle('section', parent=normal, fontSize=12, leading=20, textColor='#245447', spaceBefore=10, spaceAfter=7, keepWithNext=True)
    subheading = ParagraphStyle('subheading', parent=normal, spaceBefore=5, spaceAfter=5, keepWithNext=True)
    story = [Paragraph(escape(course), heading), Paragraph(escape(f'上课日期：{day} | 来源：{source}'), small),
             Paragraph('以下为导出资料原文，未额外生成 AI 总结。转写疑点以老师原话、课件为准。', small), Spacer(1, 14)]
    for line in body.splitlines():
        if not line.strip():
            story.append(Spacer(1, 5))
        else:
            # Chunk exceptionally long lines so no single flowable exceeds a page.
            stripped = line.strip()
            style = section if re.match(r'^- \*\*[一二三四五六七八九十]+、', stripped) or stripped.startswith('## ') else normal
            if style is normal and re.fullmatch(r'- \*\*.+\*\*', stripped):
                style = subheading
            if stripped.startswith('#'):
                stripped = stripped.lstrip('#').lstrip()
            for start in range(0, len(stripped), 700):
                markup = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', escape(stripped[start:start + 700]))
                story.append(Paragraph(markup, style))
    def footer(c, doc):
        c.saveState()
        c.setFont(font, 8)
        c.drawRightString(A4[0] - 42, 26, f'{day}  |  第 {doc.page} 页')
        c.restoreState()
    SimpleDocTemplate(str(temp), pagesize=A4, topMargin=42, bottomMargin=42, leftMargin=42, rightMargin=42).build(story, onFirstPage=footer, onLaterPages=footer)
    temp.replace(target)


def ingest(cfg, inbox, fallback, dry_run=False):
    index_path = ROOT / '归档记录' / 'index.json'
    index = read_json(index_path) if index_path.exists() else []
    known = {r['key'] for r in index}
    report = []
    for src in sorted(Path(inbox).iterdir()):
        if not src.is_file() or src.suffix.lower() not in SUPPORTED:
            continue
        try:
            day, date_method = source_day(src, fallback)
            body = extract(src)
            cid, method = identify(cfg, src.stem, body)
            if cid is None:
                report.append({'file': src.name, 'status': '待确认', 'reason': method, 'date': day})
                continue
            course = next(c for c in cfg['courses'] if c['id'] == cid)
            if not course.get('archive_enabled', True):
                report.append({'file': src.name, 'status': '未启用归档，跳过', 'course': course['name']})
                continue
            digest = hashlib.sha256(src.read_bytes()).hexdigest()
            key = f'{day}:{cid}:{digest}'
            if key in known:
                report.append({'file': src.name, 'status': '已归档，跳过重复'})
                continue
            folder = course_folder(cfg, course)
            stem = archive_stem(course, day, body)
            layout = cfg.get('archive_layout', 'lesson_folders')
            dest, original, text_path = allocate_paths(folder, day, stem, src.suffix.lower(), layout)
            warnings = []
            eligible = candidates(cfg, day)
            if cid not in {l['course'] for l in eligible['lessons']}:
                warnings.append('与候选课表不同，请核对补课/文件日期')
            if eligible['week'] is None:
                warnings.append('教学周待确认')
            if src.suffix.lower() == '.docx':
                warnings.append('生成 PDF 为文本文档；图片、公式和复杂版式请查看保留的 DOCX 原件')
            record = {'key': key, 'date': day, 'course': cid, 'course_name': course['name'], 'source': src.name,
                      'original': str(original), 'text': str(text_path), 'pdf': str(dest), 'sha256': digest,
                      'match_method': method, 'date_method': date_method, 'warnings': warnings}
            if layout == 'lesson_folders':
                record['lesson_folder'] = str(dest.parent)
            stamp = recording_start(day, body)
            if stamp:
                record['recording_start'] = stamp
            if not dry_run:
                if any(p.exists() for p in (dest, original, text_path)):
                    raise ValueError('目标文件已存在但索引未记录，保留现有文件，请人工核对')
                if src.suffix.lower() == '.pdf':
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(src, dest)
                else:
                    write_pdf(dest, course['name'], day, original.name, body)
                original.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, original)
                text_path.write_text(body, encoding='utf-8')
                index.append(record)
                save_json(index_path, index)
            known.add(key)
            report.append({'file': src.name, 'status': '预览' if dry_run else '已归档', 'course': course['name'], 'pdf': str(dest), 'warnings': warnings})
        except Exception as exc:
            report.append({'file': src.name, 'status': '失败', 'reason': str(exc)})
    if not dry_run:
        save_json(ROOT / '归档记录' / f'report-{datetime.now().strftime("%Y%m%d-%H%M%S-%f")}.json', report)
    return report


def month_bundle(cfg, month):
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', month):
        raise ValueError('月份应为 YYYY-MM')
    index_path = ROOT / '归档记录' / 'index.json'
    records = read_json(index_path) if index_path.exists() else []
    output = ROOT / '月度材料' / month
    result = []
    for course in cfg['courses']:
        if not course.get('archive_enabled', True):
            continue
        rows = [r for r in records if r['course'] == course['id'] and r['date'].startswith(month)]
        if not rows:
            result.append({'course': course['name'], 'status': '本月暂无已归档资料'})
            continue
        text = [f'# {course["name"]} · {month} 复习输入材料', '', '这是材料包，尚未生成考点与试题。', '']
        for row in sorted(rows, key=lambda r: (r['date'], r.get('recording_start', ''), r['source'])):
            source = Path(row['text'])
            if not source.is_file():
                raise ValueError(f'原文缺失：{source}')
            text += [f'## {row.get("recording_start", row["date"])} | {Path(row["pdf"]).name}', f'归档 PDF：{row["pdf"]}', '', source.read_text(encoding='utf-8'), '']
            if row.get('transcript'):
                transcript = Path(row['transcript'])
                if not transcript.is_file():
                    raise ValueError(f'逐字稿缺失：{transcript}')
                text += ['### 原始逐字稿（核对 AI 总结与老师原话，忽略无关闲聊）', f'来源链接：{row.get("feishu_url", "")}', '', transcript.read_text(encoding='utf-8'), '']
        text += ['## 给 Codex 的复习任务', '',
                 '仅依据以上课堂资料与另行提供的课件，整理核心概念、易混点及老师明确考试提示。将推测考点标为推测，标记疑似转写错误。',
                 '每门课恰好出 15 道测试题，适合该课考核内容。先考点，再题目，最后集中列出答案与解析；每题注明来源日期和文件。',
                 f'生成并检查 PDF，保存到 {Path(cfg["desktop"]) / course["folder"] / "月度复习"}。资料不足时说明缺口，不编造。']
        output.mkdir(parents=True, exist_ok=True)
        dest = output / f'{safe_name(course["name"])}_材料与复习任务.md'
        dest.write_text('\n'.join(text), encoding='utf-8')
        result.append({'course': course['name'], 'status': '已准备材料，等待 Codex 生成复习 PDF', 'path': str(dest)})
    return result


def main():
    parser = argparse.ArgumentParser(description='飞书课程助手：本地导出资料归档')
    parser.add_argument('--config', default=str(CONFIG))
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('today'); p.add_argument('--date', default=today())
    p = sub.add_parser('archive'); p.add_argument('--date'); p.add_argument('--inbox', default=str(ROOT / '收件箱')); p.add_argument('--dry-run', action='store_true')
    p = sub.add_parser('month'); p.add_argument('--month', required=True)
    args = parser.parse_args()
    cfg = read_json(args.config)
    if args.command == 'today':
        result = candidates(cfg, args.date)
        course_map = {c['id']: c for c in cfg['courses']}
        for lesson in result['lessons']:
            lesson['name'] = course_map[lesson['course']]['name']
            lesson['folder'] = str(Path(cfg['desktop']) / course_map[lesson['course']]['folder'])
        result['status'] = '仅为候选课程；教学周待确认' if result['week'] is None else '课表计划，实际调课需核对'
    elif args.command == 'archive':
        Path(args.inbox).mkdir(parents=True, exist_ok=True)
        result = ingest(cfg, args.inbox, args.date, args.dry_run)
        if not result:
            result = {'status': '收件箱没有支持的导出文件，未生成课堂 PDF'}
    else:
        result = month_bundle(cfg, args.month)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
